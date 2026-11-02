"""Exporte toutes les traductions dans docs/relecture-traductions.xlsx (une feuille par langue),
pour une relecture par des natifs. Le français sert de référence.

    python3 scripts/export_translations.py
"""
import json
import os
import re
import subprocess
from urllib.parse import unquote

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import build_langs as bl

ROOT = bl.ROOT
LANG_NAMES = {'de': 'Deutsch', 'en': 'English', 'pt': 'Português', 'lb': 'Lëtzebuergesch'}
SECTIONS = [('meta', 'Google (titre, description)'), ('skip', 'Navigation'), ('banner', 'Bandeau'), ('brand', 'Navigation'),
            ('nav', 'Navigation'), ('menu', 'Navigation'), ('lang', 'Navigation'), ('hero', 'Haut de page'),
            ('mq', 'Bandeau défilant'), ('marquee', 'Bandeau défilant'), ('svc', 'Services'), ('stmt', 'Manifeste'),
            ('why', 'Pourquoi nous'), ('phone', 'Pourquoi nous'), ('work', 'Réalisations'), ('proc', 'Méthode'),
            ('price', 'Offres'), ('faq', 'FAQ'), ('cta', 'Contact'), ('contact', 'Contact'), ('copy', 'Contact'),
            ('copied', 'Contact'), ('selected', 'Contact'), ('book', 'Contact'), ('wa', 'Contact'), ('bar', 'Barre mobile'),
            ('foot', 'Footer'), ('legal', 'Pages légales'), ('dv', 'Page devis'), ('e404', 'Page 404'), ('code', 'Animation')]


def section(key):
    for prefix, name in SECTIONS:
        if key == prefix or key.startswith(prefix + '.'):
            return name
    return 'Divers'


def readable(v):
    v = str(v)
    if 'wa.me/' in v or v.startswith('mailto:'):
        m = re.search(r'[?&](?:text|subject)=([^&]+)', v)
        return '[lien] ' + unquote(m.group(1)) if m else v
    v = re.sub(r'<em>(.*?)</em>', r'«\1»', v)
    v = re.sub(r'<[^>]+>', '', v).replace('&amp;', '&').replace('&nbsp;', ' ').replace(' ', ' ')
    return re.sub(r'\s+', ' ', v).strip()


def french_texts():
    fr = {}
    start_re = re.compile(r'<([a-zA-Z][a-zA-Z0-9]*)\b[^>]*\sdata-i18n="([^"]+)"[^>]*>')
    attr_re = re.compile(r'<[a-zA-Z][^>]*\sdata-i18n-attr="([^"]+)"[^>]*>')
    for f in ['index.html', 'devis.html', 'mentions-legales.html', '404.html']:
        html = open(os.path.join(ROOT, f), encoding='utf-8').read()
        for m in start_re.finditer(html):
            key = m.group(2)
            if key not in fr:
                end = bl.element_end(html, m.group(1), m.end())
                fr[key] = html[m.end():end]
        for m in attr_re.finditer(html):
            for pair in m.group(1).split(';'):
                attr, _, key = pair.partition(':')
                val = re.search(r'\s' + re.escape(attr.strip()) + r'="([^"]*)"', m.group(0))
                if key.strip() not in fr and val:
                    fr[key.strip()] = val.group(1)
    # Textes gérés par le JavaScript
    js = open(os.path.join(ROOT, 'assets/js/main.js'), encoding='utf-8').read() + open(os.path.join(ROOT, 'assets/js/devis.js'), encoding='utf-8').read()
    for key, val in re.findall(r"'([a-z0-9.]+)':\s*'((?:[^'\\]|\\.)*)'", js):
        fr.setdefault(key, val.replace("\\'", "'"))
    for key in ('copy', 'copied', 'selected'):
        m = re.search(key + r":\s*'([^']*)'", js)
        if m:
            fr.setdefault(key, m.group(1))
    return fr


def email_texts():
    php = open(os.path.join(ROOT, 'api/devis.php'), encoding='utf-8').read()
    block = php.split('    $T = [', 1)[1].split('\n    ];\n', 1)[0]
    out = subprocess.run(['php', '-r', '$T = [' + block + '\n]; echo json_encode($T, JSON_UNESCAPED_UNICODE);'],
                         capture_output=True, text=True, check=True).stdout
    data = json.loads(out)
    rows = {}
    for lang, d in data.items():
        flat = {}
        for k, v in d.items():
            if isinstance(v, dict):
                for k2, v2 in v.items():
                    flat[f'email.{k}.{k2}'] = v2
            else:
                flat[f'email.{k}'] = v
        rows[lang] = flat
    return rows


def main():
    fr = french_texts()
    emails = email_texts()
    hint_js = open(os.path.join(ROOT, 'assets/js/main.js'), encoding='utf-8').read()
    hints = {l: re.search(l + r": \['([^']*)', '([^']*)', '([^']*)'\]", hint_js).groups() for l in ['de', 'en', 'pt', 'lb']}

    wb = Workbook()
    info = wb.active
    info.title = "Mode d'emploi"
    lines = [
        ('DYRO WEB · Relecture des traductions du site', True),
        ('', False),
        ('Une feuille par langue : Deutsch, English, Português, Lëtzebuergesch.', False),
        ('Colonne « Français » : texte de référence. Colonne « Texte actuel » : traduction en ligne.', False),
        ('Si un texte doit changer, écrivez la version corrigée dans « Correction proposée » (sinon laissez vide).', False),
        ('Gardez les [[crochets]] tels quels : ils marquent le mot souligné dans le grand titre.', False),
        ('Les lignes « [lien] » sont le message pré-rempli d’un lien WhatsApp ou email.', False),
        ('Les lignes « email.… » sont l’email de confirmation envoyé automatiquement après une demande de devis.', False),
        ('', False),
        ('EN · One sheet per language. Write your corrected version in « Correction proposée » only where needed.', False),
        ('DE · Ein Blatt pro Sprache. Korrekturen bitte nur bei Bedarf in « Correction proposée » eintragen.', False),
        ('PT · Uma folha por língua. Escreva a correção em « Correction proposée » apenas quando necessário.', False),
        ('LB · Ee Blat pro Sprooch. Schreift Är Korrektur nëmmen am Fall vun Néierens an « Correction proposée ».', False),
    ]
    for i, (t, bold) in enumerate(lines, 1):
        c = info.cell(row=i, column=1, value=t)
        c.font = Font(bold=bold, size=14 if bold else 11)
    info.column_dimensions['A'].width = 110

    head_fill = PatternFill('solid', fgColor='0A0F1C')
    edit_fill = PatternFill('solid', fgColor='FFF7D6')
    keys = [k for k in bl.DICT['de'].keys()]
    for lang in ['de', 'en', 'pt', 'lb']:
        ws = wb.create_sheet(LANG_NAMES[lang])
        header = ['Section', 'Clé', 'Français (référence)', 'Texte actuel', 'Correction proposée', 'Commentaire']
        ws.append(header)
        for c in ws[1]:
            c.font = Font(bold=True, color='FFFFFF')
            c.fill = head_fill
            c.alignment = Alignment(vertical='center')
        rows = []
        for k in keys:
            if k not in bl.DICT[lang]:
                continue
            rows.append((section(k), k, readable(fr.get(k, '')), readable(bl.DICT[lang][k])))
        labels = ('texte', 'bouton', 'fermer')
        fr_hint = ('Ce site existe aussi en … (suggestion de langue)', 'Voir en …', 'Fermer')
        for i in range(3):
            rows.append(('Suggestion de langue', f'hint.{labels[i]}', fr_hint[i], hints[lang][i]))
        fr_mail = emails['fr']
        for k, v in emails[lang].items():
            rows.append(('Email de confirmation', k, fr_mail.get(k, ''), v))
        for r in rows:
            ws.append(list(r) + ['', ''])
        widths = [22, 20, 55, 55, 45, 30]
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical='top')
            row[4].fill = edit_fill
            row[5].fill = edit_fill
        ws.freeze_panes = 'C2'
        ws.auto_filter.ref = f'A1:F{ws.max_row}'
    out = os.path.join(ROOT, 'docs', 'relecture-traductions.xlsx')
    wb.save(out)
    print(out, {LANG_NAMES[l]: wb[LANG_NAMES[l]].max_row - 1 for l in ['de', 'en', 'pt', 'lb']})


if __name__ == '__main__':
    main()
